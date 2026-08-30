import asyncio
import io
import math
import os
import tempfile
import threading
from collections.abc import Iterator
from contextlib import contextmanager

import cv2
import pypdfium2 as pdfium
from docx import Document
from fastapi import FastAPI, File, HTTPException, Query, UploadFile
from nudenet import NudeDetector
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SOURCE_REPOSITORY = "https://github.com/cmiic/media-screener-api"


def source_url(revision: str) -> str:
    return f"{SOURCE_REPOSITORY}/tree/{revision}" if revision else SOURCE_REPOSITORY


SOURCE_URL = source_url(os.environ.get("SOURCE_REVISION", ""))

app = FastAPI(
    title="Media Screener API",
    description=f"Licensed under AGPL-3.0-or-later. [Source code]({SOURCE_URL})",
    license_info={
        "name": "GNU Affero General Public License v3.0 or later",
        "url": "https://www.gnu.org/licenses/agpl-3.0.html",
    },
)
detector = NudeDetector(model_path=os.environ.get("MODEL_PATH", "/models/640m.onnx"), inference_resolution=640)

UNSAFE_CLASSES = {
    "FEMALE_BREAST_EXPOSED",
    "FEMALE_GENITALIA_EXPOSED",
    "MALE_GENITALIA_EXPOSED",
    "BUTTOCKS_EXPOSED",
    "ANUS_EXPOSED",
}

# Supported MIME types
IMAGE_TYPES = {"image/jpeg", "image/png", "image/webp", "image/gif", "image/bmp"}
VIDEO_TYPES = {"video/mp4", "video/webm", "video/avi", "video/quicktime", "video/x-msvideo", "video/x-matroska"}
PDF_TYPES = {"application/pdf"}
DOCX_TYPES = {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
PPTX_TYPES = {"application/vnd.openxmlformats-officedocument.presentationml.presentation"}
PDF_RENDER_SCALE = 150 / 72
MAX_UPLOAD_SIZE = int(os.environ.get("MAX_UPLOAD_SIZE", str(200 * 1024 * 1024)))
MAX_PDF_PAGES = int(os.environ.get("MAX_PDF_PAGES", "100"))
MAX_PDF_RENDERED_PIXELS = int(os.environ.get("MAX_PDF_RENDERED_PIXELS", "100000000"))
MAX_VIDEO_DURATION_SECONDS = float(os.environ.get("MAX_VIDEO_DURATION_SECONDS", "600"))
MAX_VIDEO_SAMPLED_FRAMES = int(os.environ.get("MAX_VIDEO_SAMPLED_FRAMES", "120"))
MAX_DOCUMENT_IMAGES = int(os.environ.get("MAX_DOCUMENT_IMAGES", "100"))
MAX_CONCURRENT_SCANS = int(os.environ.get("MAX_CONCURRENT_SCANS", "1"))
SCAN_QUEUE_TIMEOUT_SECONDS = float(os.environ.get("SCAN_QUEUE_TIMEOUT_SECONDS", "5"))
UPLOAD_TIMEOUT_SECONDS = float(os.environ.get("UPLOAD_TIMEOUT_SECONDS", "10"))
PROCESSING_TIMEOUT_SECONDS = float(os.environ.get("PROCESSING_TIMEOUT_SECONDS", "40"))

pdfium_lock = threading.Lock()
scan_slots = asyncio.Semaphore(MAX_CONCURRENT_SCANS)

# File extension fallback mapping
EXT_TO_TYPE = {
    ".jpg": "image", ".jpeg": "image", ".png": "image", ".webp": "image", ".gif": "image", ".bmp": "image",
    ".mp4": "video", ".webm": "video", ".avi": "video", ".mov": "video", ".mkv": "video",
    ".pdf": "pdf",
    ".docx": "docx",
    ".pptx": "pptx",
}


@app.get("/health")
async def health() -> dict[str, str]:
    """Simple readiness endpoint for container health checks."""
    return {"status": "healthy"}


@app.get("/source")
async def source() -> dict[str, str]:
    """Provide the Corresponding Source location required by the AGPL."""
    return {"license": "AGPL-3.0-or-later", "source": SOURCE_URL}


async def read_upload(file: UploadFile) -> bytes:
    """Read an upload without allowing the application buffer to exceed its limit."""
    if file.size is not None and file.size > MAX_UPLOAD_SIZE:
        raise HTTPException(status_code=413, detail=f"File exceeds the {MAX_UPLOAD_SIZE}-byte upload limit")

    contents = await file.read(MAX_UPLOAD_SIZE + 1)
    if len(contents) > MAX_UPLOAD_SIZE:
        raise HTTPException(status_code=413, detail=f"File exceeds the {MAX_UPLOAD_SIZE}-byte upload limit")
    return contents


def release_scan_slot(worker: asyncio.Task, slots: asyncio.Semaphore) -> None:
    """Release a slot after timed-out worker code has actually stopped."""
    if not worker.cancelled():
        worker.exception()
    slots.release()


async def process_upload(file: UploadFile, operation, *args):
    """Read and process one upload within queue, time, and concurrency bounds."""
    slots = scan_slots
    try:
        await asyncio.wait_for(slots.acquire(), timeout=SCAN_QUEUE_TIMEOUT_SECONDS)
    except TimeoutError as exc:
        raise HTTPException(status_code=503, detail="Scanner is at capacity; retry later") from exc

    release_slot = True
    try:
        try:
            contents = await asyncio.wait_for(read_upload(file), timeout=UPLOAD_TIMEOUT_SECONDS)
        except TimeoutError as exc:
            raise HTTPException(status_code=408, detail="Upload timed out") from exc

        worker = asyncio.create_task(asyncio.to_thread(operation, contents, *args))
        try:
            return await asyncio.wait_for(asyncio.shield(worker), timeout=PROCESSING_TIMEOUT_SECONDS)
        except TimeoutError as exc:
            if worker.done():
                raise
            release_slot = False
            worker.add_done_callback(lambda completed: release_scan_slot(completed, slots))
            raise HTTPException(status_code=504, detail="Media processing timed out") from exc
        except asyncio.CancelledError:
            release_slot = False
            worker.add_done_callback(lambda completed: release_scan_slot(completed, slots))
            raise
    finally:
        if release_slot:
            slots.release()


def detect_file_type(content_type: str, filename: str) -> str:
    """Detect file type from MIME type or extension."""
    if content_type in IMAGE_TYPES:
        return "image"
    if content_type in VIDEO_TYPES:
        return "video"
    if content_type in PDF_TYPES:
        return "pdf"
    if content_type in DOCX_TYPES:
        return "docx"
    if content_type in PPTX_TYPES:
        return "pptx"

    # Fallback to extension
    if filename:
        ext = "." + filename.lower().split(".")[-1] if "." in filename else ""
        return EXT_TO_TYPE.get(ext, "unknown")

    return "unknown"


def extract_video_frames(video_bytes: bytes, sample_interval: float = 2.0) -> list[tuple[bytes, float]]:
    """Extract frames from video at given interval. Returns list of (frame_bytes, timestamp)."""
    frames = []

    with tempfile.NamedTemporaryFile(suffix=".mp4", delete=True) as tmp:
        tmp.write(video_bytes)
        tmp.flush()

        cap = cv2.VideoCapture(tmp.name)
        fps = cap.get(cv2.CAP_PROP_FPS) or 30
        frame_count = cap.get(cv2.CAP_PROP_FRAME_COUNT)
        duration = frame_count / fps if frame_count > 0 else 0
        if duration > MAX_VIDEO_DURATION_SECONDS:
            cap.release()
            raise HTTPException(
                status_code=422,
                detail=f"Video exceeds the {MAX_VIDEO_DURATION_SECONDS:g}-second duration limit",
            )

        frame_interval = max(1, int(fps * sample_interval))
        frame_idx = 0

        try:
            while True:
                ret, frame = cap.read()
                if not ret:
                    break

                timestamp = frame_idx / fps
                if timestamp > MAX_VIDEO_DURATION_SECONDS:
                    raise HTTPException(
                        status_code=422,
                        detail=f"Video exceeds the {MAX_VIDEO_DURATION_SECONDS:g}-second duration limit",
                    )

                if frame_idx % frame_interval == 0:
                    if len(frames) >= MAX_VIDEO_SAMPLED_FRAMES:
                        raise HTTPException(
                            status_code=422,
                            detail=f"Video exceeds the {MAX_VIDEO_SAMPLED_FRAMES}-sample frame limit",
                        )
                    _, buffer = cv2.imencode(".jpg", frame)
                    frames.append((buffer.tobytes(), timestamp))

                frame_idx += 1
        finally:
            cap.release()

    return frames


@contextmanager
def render_pdf_pages(pdf_bytes: bytes) -> Iterator[tuple[int, Iterator[tuple[bytes, int]]]]:
    """Yield 150 DPI PDF page renders without retaining them in memory."""
    with pdfium_lock:
        document = pdfium.PdfDocument(pdf_bytes)
        total_pages = len(document)

    try:
        if total_pages > MAX_PDF_PAGES:
            raise HTTPException(status_code=422, detail=f"PDF exceeds the {MAX_PDF_PAGES}-page limit")

        def pages() -> Iterator[tuple[bytes, int]]:
            rendered_pixels = 0
            for page_index in range(total_pages):
                with pdfium_lock:
                    page = document[page_index]
                    try:
                        width, height = page.get_size()
                        rendered_pixels += math.ceil(width * PDF_RENDER_SCALE) * math.ceil(height * PDF_RENDER_SCALE)
                        if rendered_pixels > MAX_PDF_RENDERED_PIXELS:
                            raise HTTPException(
                                status_code=422,
                                detail=f"PDF exceeds the {MAX_PDF_RENDERED_PIXELS}-pixel render limit",
                            )
                        bitmap = page.render(scale=PDF_RENDER_SCALE)
                        try:
                            image = bitmap.to_pil()
                            try:
                                output = io.BytesIO()
                                image.save(output, format="PNG")
                                image_bytes = output.getvalue()
                            finally:
                                image.close()
                        finally:
                            bitmap.close()
                    finally:
                        page.close()

                yield image_bytes, page_index + 1

        yield total_pages, pages()
    finally:
        with pdfium_lock:
            document.close()


def extract_docx_images(docx_bytes: bytes) -> list[tuple[bytes, str]]:
    """Extract embedded images from DOCX. Returns list of (image_bytes, image_name)."""
    images = []

    doc = Document(io.BytesIO(docx_bytes))
    for rel in doc.part.rels.values():
        if "image" in rel.target_ref:
            if len(images) >= MAX_DOCUMENT_IMAGES:
                raise HTTPException(
                    status_code=422,
                    detail=f"Document exceeds the {MAX_DOCUMENT_IMAGES}-image limit",
                )
            try:
                img_bytes = rel.target_part.blob
                images.append((img_bytes, rel.target_ref))
            except Exception:
                pass

    return images


def extract_pptx_images(pptx_bytes: bytes) -> list[tuple[bytes, str]]:
    """Extract embedded images from PPTX. Returns list of (image_bytes, location)."""
    images = []

    prs = Presentation(io.BytesIO(pptx_bytes))
    for slide_num, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                if len(images) >= MAX_DOCUMENT_IMAGES:
                    raise HTTPException(
                        status_code=422,
                        detail=f"Document exceeds the {MAX_DOCUMENT_IMAGES}-image limit",
                    )
                try:
                    img_bytes = shape.image.blob
                    images.append((img_bytes, f"slide_{slide_num}"))
                except Exception:
                    pass

    return images


def classify_image(image_bytes: bytes, threshold: float) -> dict:
    """Classify a single image and return detection info."""
    detections = detector.detect(image_bytes)
    unsafe_detections = [d for d in detections if d["class"] in UNSAFE_CLASSES and d["score"] >= threshold]

    return {
        "unsafe": len(unsafe_detections) > 0,
        "confidence": max((d["score"] for d in unsafe_detections), default=0.0),
        "detected_classes": [d["class"] for d in unsafe_detections],
        "all_detections": detections,
    }


def infer_contents(contents: bytes, content_type: str, filename: str, sample_interval: float):
    """Run detailed inference after upload bounds have been enforced."""
    file_type = detect_file_type(content_type, filename)

    if file_type == "image":
        detections = detector.detect(contents)
        return {"format": file_type, "prediction": detections}

    elif file_type == "video":
        frames = extract_video_frames(contents, sample_interval)
        results = []
        for frame_bytes, timestamp in frames:
            detections = detector.detect(frame_bytes)
            if detections:
                results.append({"timestamp": round(timestamp, 2), "detections": detections})
        return {"format": file_type, "frames_checked": len(frames), "predictions": results}

    elif file_type == "pdf":
        results = []
        with render_pdf_pages(contents) as (total_pages, pages):
            for img_bytes, page_num in pages:
                detections = detector.detect(img_bytes)
                if detections:
                    results.append({"page": page_num, "detections": detections})
        return {"format": file_type, "pages_checked": total_pages, "predictions": results}

    elif file_type == "docx":
        images = extract_docx_images(contents)
        results = []
        for img_bytes, img_name in images:
            detections = detector.detect(img_bytes)
            if detections:
                results.append({"image": img_name, "detections": detections})
        return {"format": file_type, "images_checked": len(images), "predictions": results}

    elif file_type == "pptx":
        images = extract_pptx_images(contents)
        results = []
        for img_bytes, location in images:
            detections = detector.detect(img_bytes)
            if detections:
                results.append({"location": location, "detections": detections})
        return {"format": file_type, "images_checked": len(images), "predictions": results}

    else:
        return {"error": f"Unsupported file type: {content_type}", "filename": filename}


@app.post("/infer")
async def infer(
    file: UploadFile = File(...),
    sample_interval: float = Query(5.0, ge=0.1, le=60, description="Seconds between video frame samples"),
):
    """Detect nudity in an uploaded file (image, video, PDF, DOCX, PPTX)."""
    return await process_upload(file, infer_contents, file.content_type or "", file.filename or "", sample_interval)


def classify_contents(
    contents: bytes,
    content_type: str,
    filename: str,
    threshold: float,
    sample_interval: float,
    early_exit: bool,
):
    """Classify an upload after its request-level bounds have been enforced."""
    file_type = detect_file_type(content_type, filename)

    if file_type == "image":
        result = classify_image(contents, threshold)
        return {
            "format": file_type,
            "unsafe": result["unsafe"],
            "confidence": result["confidence"],
            "detected_classes": result["detected_classes"],
        }

    elif file_type == "video":
        frames = extract_video_frames(contents, sample_interval)
        all_detected_classes = []
        max_confidence = 0.0
        first_unsafe_at = None
        frames_checked = 0

        for frame_bytes, timestamp in frames:
            frames_checked += 1
            result = classify_image(frame_bytes, threshold)

            if result["unsafe"]:
                all_detected_classes.extend(result["detected_classes"])
                max_confidence = max(max_confidence, result["confidence"])
                if first_unsafe_at is None:
                    first_unsafe_at = timestamp
                if early_exit:
                    break

        return {
            "format": file_type,
            "unsafe": len(all_detected_classes) > 0,
            "confidence": max_confidence,
            "detected_classes": list(set(all_detected_classes)),
            "frames_checked": frames_checked,
            "total_frames": len(frames),
            "first_unsafe_at": round(first_unsafe_at, 2) if first_unsafe_at is not None else None,
        }

    elif file_type == "pdf":
        all_detected_classes = []
        max_confidence = 0.0
        first_unsafe_at = None
        pages_checked = 0

        with render_pdf_pages(contents) as (total_pages, pages):
            for img_bytes, page_num in pages:
                pages_checked += 1
                result = classify_image(img_bytes, threshold)

                if result["unsafe"]:
                    all_detected_classes.extend(result["detected_classes"])
                    max_confidence = max(max_confidence, result["confidence"])
                    if first_unsafe_at is None:
                        first_unsafe_at = page_num
                    if early_exit:
                        break

        return {
            "format": file_type,
            "unsafe": len(all_detected_classes) > 0,
            "confidence": max_confidence,
            "detected_classes": list(set(all_detected_classes)),
            "pages_checked": pages_checked,
            "total_pages": total_pages,
            "first_unsafe_at": first_unsafe_at,
        }

    elif file_type in ("docx", "pptx"):
        if file_type == "docx":
            images = extract_docx_images(contents)
        else:
            images = extract_pptx_images(contents)

        all_detected_classes = []
        max_confidence = 0.0
        first_unsafe_at = None
        images_checked = 0

        for img_bytes, location in images:
            images_checked += 1
            result = classify_image(img_bytes, threshold)

            if result["unsafe"]:
                all_detected_classes.extend(result["detected_classes"])
                max_confidence = max(max_confidence, result["confidence"])
                if first_unsafe_at is None:
                    first_unsafe_at = location
                if early_exit:
                    break

        return {
            "format": file_type,
            "unsafe": len(all_detected_classes) > 0,
            "confidence": max_confidence,
            "detected_classes": list(set(all_detected_classes)),
            "images_checked": images_checked,
            "total_images": len(images),
            "first_unsafe_at": first_unsafe_at,
        }

    else:
        return {"error": f"Unsupported file type: {content_type}", "filename": filename}


@app.post("/classify")
async def classify(
    file: UploadFile = File(...),
    threshold: float = Query(0.35, ge=0, le=1, description="Minimum confidence for unsafe detection"),
    sample_interval: float = Query(5.0, ge=0.1, le=60, description="Seconds between video frame samples"),
    early_exit: bool = Query(True, description="Stop on first unsafe detection"),
):
    """Classify file as safe or unsafe (image, video, PDF, DOCX, PPTX)."""
    return await process_upload(
        file,
        classify_contents,
        file.content_type or "",
        file.filename or "",
        threshold,
        sample_interval,
        early_exit,
    )

